import openpyxl

# Welcome
print("================================")
print("WELCOME TO BANK MANAGEMENT SYSTEM")
print("================================")

# Operations available
print("\nYou can perform the following operations:")
print("1. Create Account")
print("2. Deposit Money")
print("3. Withdraw Money")
print("4. Check Balance")

# Number of operations
num_op = int(input("\nEnter number of operations: "))


# Excel file
try:
    workbook = openpyxl.load_workbook("banking_system.xlsx")
    sheet = workbook.active
except:
    workbook = openpyxl.Workbook()
    sheet = workbook.active
    sheet.append(["Account Number", "Name", "Mobile Number", "PIN", "Balance"])
    workbook.save("banking_system.xlsx")


# Create Account
def create_account():
    print("\n--- Create Account ---")

    name = input("Enter Account Holder Name: ")
    mobile = input("Enter Mobile Number: ")
    pin = input("Enter 4-digit PIN: ")
    balance = float(input("Enter Initial Deposit: "))

    account_no = 1001

    if sheet.max_row > 1:
        account_no = sheet.cell(sheet.max_row, 1).value + 1

    sheet.append([account_no, name, mobile, pin, balance])
    workbook.save("banking_system.xlsx")

    print("Account Created Successfully!")
    print("Your Account Number:", account_no)


# Find Account
def find_account(account_no):
    for i in range(2, sheet.max_row + 1):
        if sheet.cell(i, 1).value == account_no:
            return i
    return 0


# Verify PIN
def verify_pin(row):
    pin = input("Enter PIN: ")

    if sheet.cell(row, 4).value == pin:
        return True
    else:
        print("Incorrect PIN!")
        return False


# Deposit Money
def deposit_money():
    print("\n--- Deposit Money ---")

    account_no = int(input("Enter Account Number: "))
    row = find_account(account_no)

    if row == 0:
        print("Account Not Found!")
    else:
        if verify_pin(row):
            amount = float(input("Enter Deposit Amount: "))

            if amount > 0:
                balance = sheet.cell(row, 5).value
                balance = balance + amount
                sheet.cell(row, 5).value = balance

                workbook.save("banking_system.xlsx")

                print("Deposit Successful!")
                print("Updated Balance:", balance)
            else:
                print("Invalid Amount!")


# Withdraw Money
def withdraw_money():
    print("\n--- Withdraw Money ---")

    account_no = int(input("Enter Account Number: "))
    row = find_account(account_no)

    if row == 0:
        print("Account Not Found!")
    else:
        if verify_pin(row):
            amount = float(input("Enter Withdrawal Amount: "))
            balance = sheet.cell(row, 5).value

            if amount <= 0:
                print("Invalid Amount!")
            elif amount > balance:
                print("Insufficient Balance!")
            else:
                balance = balance - amount
                sheet.cell(row, 5).value = balance

                workbook.save("banking_system.xlsx")

                print("Withdrawal Successful!")
                print("Updated Balance:", balance)


# Check Balance
def check_balance():
    print("\n--- Check Balance ---")

    account_no = int(input("Enter Account Number: "))
    row = find_account(account_no)

    if row == 0:
        print("Account Not Found!")
    else:
        if verify_pin(row):
            print("Account Number:", sheet.cell(row, 1).value)
            print("Account Holder:", sheet.cell(row, 2).value)
            print("Current Balance:", sheet.cell(row, 5).value)


# Operations
for i in range(1, num_op + 1, 1):

    print("\n--- Operation", i, "---")

    choice = int(input("Enter your choice: "))

    if choice == 1:
        create_account()

    elif choice == 2:
        deposit_money()

    elif choice == 3:
        withdraw_money()

    elif choice == 4:
        check_balance()

    else:
        print("Invalid Option!")


# Exit
print("\n5. Exit")

choice = int(input("Enter your choice: "))

if choice == 5:
    print("\nThank you for using Bank Management System!")
else:
    print("Invalid Option!")