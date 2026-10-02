import openpyxl

workbook = openpyxl.Workbook()
sheet = workbook.active

sheet["A1"] = "Account No"
sheet["B1"] = "Name"
sheet["C1"] = "Balance"

sheet.append([101, "Kirti", 5000])

workbook.save("bank.xlsx")

print("Data saved in Excel!")


import openpyxl

workbook = openpyxl.load_workbook("bank.xlsx")
account_no = input("Enter Account Number: ")
name = input("Enter Name: ")
balance = float(input("Enter Initial Balance: "))
sheet.append([account_no, name, balance])
workbook.save("bank.xlsx")
for row in sheet.iter_rows(values_only=True):
    print(row)